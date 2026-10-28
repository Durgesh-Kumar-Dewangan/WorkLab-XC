"""Pulls plain text out of an uploaded syllabus: .txt/.md directly, .pdf with pypdf, .docx with python-docx."""
from __future__ import annotations

import io
import re

from .errors import ApiError

_BAD = ApiError(415, "unsupported_file", "That file type is not supported. Use .txt, .md, .pdf or .docx.")


def _ext(n: str) -> str:
    m = re.search(r"\.([a-z0-9]+)$", n.lower())
    return m.group(1) if m else ""


def kind_of(name: str, data: bytes) -> str:
    e = _ext(name)
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        if e == "docx":
            return "docx"
        raise ApiError(415, "unsupported_file", "That file type is not supported. Use .txt, .md, .pdf or .docx.")
    if e in ("txt", "md", "markdown", "text", "") and b"\x00" not in data[:2048]:
        return "text"
    raise ApiError(415, "unsupported_file", "That file type is not supported. Use .txt, .md, .pdf or .docx.")


def extract_text(name: str, data: bytes) -> str:
    kind = kind_of(name, data)
    try:
        if kind == "text":
            return data.decode("utf-8", "replace").lstrip("﻿")
        if kind == "pdf":
            from pypdf import PdfReader

            r = PdfReader(io.BytesIO(data))
            if len(r.pages) > 150:
                raise ApiError(413, "too_many_pages", "That PDF has more than 150 pages. Upload just the syllabus pages.")
            return "\n".join((p.extract_text() or "") for p in r.pages)
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except ApiError:
        raise
    except Exception:  # noqa: BLE001
        raise ApiError(422, "unreadable_file", "We could not read text from that file. If it is a scanned PDF, paste the text instead.") from None


def tidy(text: str) -> str:
    t = text.replace("\r", "").replace("\x00", "")
    t = re.sub(r"[•◦▪●‣∙]\s*", "- ", t)
    t = re.sub(r"[ \t]+\n", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()
