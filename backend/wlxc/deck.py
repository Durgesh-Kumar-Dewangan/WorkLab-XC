"""Builds the 16:9 PowerPoint: a title slide, then per topic a heading slide and one slide per generated slide,
with the narration as speaker notes. Uses python-pptx."""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .render import AI_NOTE, WATERMARK
from .styles import STYLES, mix


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def build_pptx(course: dict, done: list[dict], watermark: bool, theme: dict | None = None) -> bytes:
    theme = theme or STYLES["clean"]
    p = Presentation()
    p.slide_width, p.slide_height = Inches(10), Inches(5.625)
    p.core_properties.title = course["title"][:250]
    p.core_properties.author = "WorkLab XC"
    blank = p.slide_layouts[6]
    INK, ACCENT, MUTED, FONT, HEAD = theme["ink"], theme["accent"], theme["muted"], theme["bodyFont"], theme["headFont"]
    FAINT, PANEL = mix(theme["bg"], theme["ink"], 0.07), mix(theme["bg"], theme["accent"], 0.1)

    def text(s, t, x, y, w, h, font, size, color, bold=False, align=None, valign=None, fill=None, rotate=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if valign:
            tf.vertical_anchor = valign
        r = tf.paragraphs[0].add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = font, Pt(size), bold, _rgb(color)
        if align:
            tf.paragraphs[0].alignment = align
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = _rgb(fill)
        if rotate:
            tb.rotation = rotate
        return tb

    def new_slide(bg):
        s = p.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(bg)
        if watermark:  # first, so it sits behind the text
            text(s, "WorkLab XC · free plan", 1.2, 1.7, 7.6, 2.2, FONT, 54, FAINT, True, PP_ALIGN.CENTER, rotate=330)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(0.14))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(ACCENT)
        bar.line.fill.background()
        return s

    n = [1]

    def chrome(s):
        text(s, f"{WATERMARK if watermark else 'WorkLab XC'}  ·  AI-generated, review before use", 0.5, 5.2, 8, 0.3, FONT, 9, MUTED)
        text(s, str(n[0]), 9, 5.2, 0.5, 0.3, FONT, 9, MUTED, align=PP_ALIGN.RIGHT)
        n[0] += 1

    t0 = new_slide(theme["bgAlt"])
    text(t0, course["title"], 0.6, 1.6, 8.8, 1.2, HEAD, 36, INK, True)
    text(t0, "Study deck", 0.6, 2.8, 8.8, 0.5, FONT, 18, ACCENT)
    text(t0, AI_NOTE, 0.6, 3.5, 8.8, 0.6, FONT, 12, MUTED)
    chrome(t0)

    by = {d["id"]: d for d in done}
    for tid in course["order"]:
        d = by.get(tid)
        if not d:
            continue
        topic = course["topics"][tid]
        head = new_slide(theme["bgAlt"])
        w = 5.2 if d.get("image") else 8.8
        text(head, topic["title"], 0.6, 1.5, w, 1, HEAD, 32, INK, True)
        text(head, d["draft"]["summary"], 0.6, 2.6, w, 1.4, FONT, 16, MUTED, valign=MSO_ANCHOR.TOP)
        if d.get("image"):
            try:
                head.shapes.add_picture(io.BytesIO(d["image"]["bytes"]), Inches(6.1), Inches(1.5), Inches(3.4), Inches(1.9125))
            except Exception:  # noqa: BLE001  - a broken image must not lose the whole deck
                pass
        head.notes_slide.notes_text_frame.text = "\n".join(f"• {o}" for o in d["draft"]["objectives"])
        chrome(head)
        for sl in d["draft"]["slides"]:
            s = new_slide(theme["bg"])
            text(s, sl["title"], 0.5, 0.35, 9, 0.8, HEAD, 26, INK, True)
            chars = len("".join(sl["bullets"]))
            size = 14 if chars > 600 else 16 if chars > 400 else 18
            tb = s.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(8.8), Inches(2.9 if sl.get("formula") else 3.6))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for i, b in enumerate(sl["bullets"]):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.space_after = Pt(8)
                r = para.add_run()
                r.text = f"• {b}"
                r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), _rgb(INK)
            if sl.get("formula"):
                text(s, sl["formula"], 0.6, 4.35, 8.8, 0.6, "Consolas", 18, ACCENT, True, PP_ALIGN.CENTER, fill=PANEL)
            s.notes_slide.notes_text_frame.text = sl["narration"]
            chrome(s)
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()
