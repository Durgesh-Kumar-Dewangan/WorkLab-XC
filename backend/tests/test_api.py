import json
import time
import unittest
from urllib.parse import urlencode

from wlxc import integrations as integ, jobs, llm, net
from wlxc.auth import effective_plan
from wlxc.image import generate_image
from wlxc.lesson import generate_best
from wlxc.net import Resp
from wlxc.store import get_store

from .helpers import AUTH_CRON, Base, PNG, STATS, call, cookie_of, delete, get, hmac_hex, post, setup


class ApiBase(Base):
    def make_course(self, cookie):
        r = post("/api/courses", {"text": STATS}, cookie=cookie)
        self.assertEqual(r.status, 201, r.text)
        return r.json()

    def run_done(self, cookie, id):
        jobs.run_job(id, 120_000)
        return get(f"/api/jobs/{id}", cookie=cookie).json()

    def start_job(self, cookie, course, n=1, start=0, **extra):
        r = post("/api/jobs", {"courseId": course["id"], "topicIds": course["course"]["order"][start:start + n], **extra}, cookie=cookie)
        return r


class TestAccounts(ApiBase):
    def test_register_sets_httponly_cookie_and_reads_account_back(self):
        cookie, res = self.sign_up()
        self.assertEqual(res.status, 201)
        self.assertIn("HttpOnly", res.headers["set-cookie"])
        self.assertIn("SameSite=Lax", res.headers["set-cookie"])
        j = get("/api/me", cookie=cookie).json()
        self.assertEqual(j["user"]["email"], "asha@example.com")
        self.assertEqual(j["user"]["plan"], "prototype")
        self.assertFalse(j["feedbackRequired"])

    def test_never_stores_or_returns_the_password(self):
        _, res = self.sign_up()
        self.assertNotIn("hash", res.text)
        self.assertNotIn("S3cure", res.text)
        self.assertTrue(self.user_doc("asha@example.com")["hash"].startswith("scrypt$"))

    def test_rejects_duplicates_bad_input_and_wrong_passwords(self):
        self.sign_up()
        self.assertEqual(self.sign_up()[1].status, 409)
        self.assertEqual(post("/api/auth/register", {"name": "x", "email": "nope", "password": "S3cure-pass!"}).status, 422)
        self.assertEqual(post("/api/auth/register", {"name": "x", "email": "b@example.com", "password": "short"}).status, 422)
        self.assertEqual(post("/api/auth/login", {"email": "asha@example.com", "password": "wrong-pass"}).status, 401)
        self.assertEqual(post("/api/auth/login", {"email": "ghost@example.com", "password": "whatever1"}).status, 401)
        self.assertEqual(post("/api/auth/login", {"email": "asha@example.com", "password": "S3cure-pass!"}).status, 200)

    def test_blocks_cross_site_posts_and_unauthenticated_calls(self):
        r = post("/api/auth/register", {"name": "x", "email": "c@example.com", "password": "S3cure-pass!"}, origin="https://evil.example")
        self.assertEqual(r.status, 403)
        self.assertEqual(get("/api/me").status, 401)
        self.assertEqual(get("/api/jobs").status, 401)

    def test_rate_limits_sign_in_attempts(self):
        last = 0
        for i in range(12):
            last = post("/api/auth/login", {"email": "asha@example.com", "password": f"wrong-{i}"}).status
        self.assertEqual(last, 429)

    def test_uniform_json_errors_with_request_id(self):
        r = post("/api/auth/register", raw="{not json")
        self.assertEqual(r.status, 400)
        self.assertEqual(r.json()["error"], "bad_json")
        self.assertTrue(r.json()["requestId"])


class TestCourses(ApiBase):
    def test_parses_adds_prerequisites_without_cycles_and_stores(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        order = c["course"]["order"]
        self.assertEqual(len(order), 12)
        self.assertTrue(c["aiPrerequisites"])
        self.assertEqual(len(c["course"]["topics"][order[0]].get("inf") or []), 0)  # the cycle edge back to topic 1 was dropped
        self.assertEqual(c["course"]["topics"][order[1]]["inf"][0], order[0])
        self.assertEqual(get("/api/courses", cookie=cookie).json()["courses"][0]["id"], c["id"])
        self.assertEqual(get(f"/api/courses/{c['id']}", cookie=cookie).status, 200)

    def test_courses_are_private(self):
        a, _ = self.sign_up("a@example.com")
        b, _ = self.sign_up("b@example.com")
        c = self.make_course(a)
        self.assertEqual(get(f"/api/courses/{c['id']}", cookie=b).status, 403)

    def test_works_when_every_ai_provider_is_down(self):
        self.w.fail |= {"az.test", "api.cloudflare.com"}
        cookie, _ = self.sign_up()
        self.assertFalse(self.make_course(cookie)["aiPrerequisites"])

    def test_md_upload_ok_pdf_refused_on_free_plan(self):
        cookie, _ = self.sign_up()
        ok = call("POST", "/api/courses", cookie=cookie, files={"file": ("stats.md", STATS.encode(), "text/markdown")})
        self.assertEqual(ok.status, 201, ok.text)
        no = call("POST", "/api/courses", cookie=cookie, files={"file": ("x.pdf", b"%PDF-\x01\x02\x03", "application/pdf")})
        self.assertEqual(no.status, 402)

    def test_rejects_empty_syllabus(self):
        cookie, _ = self.sign_up()
        self.assertEqual(post("/api/courses", {"text": "short"}, cookie=cookie).status, 422)

    def test_real_pdf_and_docx_are_read_on_paid_plan(self):
        import io
        import docx
        from pypdf import PdfWriter
        cookie, _ = self.sign_up()
        self.upgrade("asha@example.com")
        d = docx.Document()
        for line in STATS.split("\n"):
            d.add_paragraph(line)
        buf = io.BytesIO()
        d.save(buf)
        r = call("POST", "/api/courses", cookie=cookie, files={"file": ("syllabus.docx", buf.getvalue(), "application/octet-stream")})
        self.assertEqual(r.status, 201, r.text)
        self.assertEqual(len(r.json()["course"]["order"]), 12)
        pdf = io.BytesIO()
        w = PdfWriter()
        w.add_blank_page(200, 200)
        w.write(pdf)
        r2 = call("POST", "/api/courses", cookie=cookie, files={"file": ("blank.pdf", pdf.getvalue(), "application/pdf")})
        self.assertEqual(r2.status, 422)  # a blank (scanned-like) PDF has no text: asks to paste it instead


class TestGeneration(ApiBase):
    def test_runs_to_completion_with_notes_deck_and_course_file(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        create = self.start_job(cookie, c, 3)
        self.assertEqual(create.status, 202, create.text)
        job = create.json()["job"]
        done = self.run_done(cookie, job["id"])["job"]
        self.assertEqual(done["status"], "done")
        self.assertEqual(done["progress"]["done"], 3)
        self.assertEqual(sorted(done["downloads"]), ["json", "md", "pptx"])
        pptx = get(f"/api/jobs/{job['id']}/download?format=pptx", cookie=cookie)
        self.assertEqual(pptx.status, 200)
        self.assertIn("presentationml", pptx.headers["content-type"])
        self.assertIn('attachment; filename="Introductory-Statistics.pptx"', pptx.headers["content-disposition"])
        self.assertEqual(pptx.body[:2], b"PK")
        self.assertGreater(len(pptx.body), 20_000)
        # it is a real deck: re-open it and count slides
        import io
        from pptx import Presentation
        prs = Presentation(io.BytesIO(pptx.body))
        self.assertGreater(len(prs.slides), 3 * 6)
        self.assertTrue(any(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides))
        md = get(f"/api/jobs/{job['id']}/download?format=md", cookie=cookie).text
        self.assertIn("AI-generated study material", md)
        self.assertIn("Check yourself", md)
        self.assertIn("free plan", md)
        js = get(f"/api/jobs/{job['id']}/download?format=json", cookie=cookie).json()
        t0 = js["topics"][js["order"][0]]
        self.assertIn("explained simply", t0["summary"])
        self.assertGreater(len(t0["slides"]), 4)
        self.assertEqual(t0["qs"][0]["type"], "mcq")
        self.assertEqual(js["ai"]["models"][0], "gpt-4o-mini")
        self.assertEqual(js["topics"][js["order"][5]]["summary"], "")  # a topic nobody generated holds no placeholder text
        stored = get(f"/api/courses/{c['id']}", cookie=cookie).json()
        self.assertIn("explained simply", stored["course"]["topics"][c["course"]["order"][0]]["summary"])

    def test_paid_plan_removes_watermark_and_allows_images(self):
        cookie, _ = self.sign_up("paid@example.com")
        self.upgrade("paid@example.com")
        c = self.make_course(cookie)
        job = self.start_job(cookie, c, 2, images=True, instructions="use cricket examples").json()["job"]
        self.assertEqual(self.run_done(cookie, job["id"])["job"]["status"], "done")
        md = get(f"/api/jobs/{job['id']}/download?format=md", cookie=cookie).text
        self.assertNotIn("free plan", md)
        j = get_store().get("jobs", job["id"])
        self.assertEqual(j["usage"]["images"], 2)
        self.assertEqual(j["opts"]["candidates"], 2)
        self.assertTrue(all(t.get("imageBlob") for t in j["topics"].values()))
        prompts = " ".join(x["body"] for x in self.w.calls if "az.test" in x["url"])
        self.assertIn("use cricket examples", prompts)

    def test_enforces_plan_limits_daily_quota_and_instruction_gating(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        mk = lambda start, n, **e: self.start_job(cookie, c, n, start, **e)  # noqa: E731
        self.assertEqual(mk(0, 6).status, 402)
        self.assertEqual(mk(0, 2, images=True).status, 402)
        self.assertEqual(mk(0, 2, instructions="be funny").status, 402)
        self.assertEqual(post("/api/jobs", {"courseId": c["id"], "topicIds": ["nope"]}, cookie=cookie).status, 422)
        self.assertEqual([mk(0, 1).status, mk(1, 1).status], [202, 202])
        third = mk(2, 1)
        self.assertEqual(third.status, 429)
        self.assertEqual(third.json()["error"], "quota_exceeded")

    def test_falls_back_to_next_provider(self):
        self.w.fail.add("az.test")
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        job = self.start_job(cookie, c, 1).json()["job"]
        self.assertEqual(self.run_done(cookie, job["id"])["job"]["status"], "done")
        j = get_store().get("jobs", job["id"])
        self.assertEqual(list(j["topics"].values())[0]["provider"], "cloudflare")

    def test_fails_cleanly_and_refunds_when_every_provider_is_down(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        job = self.start_job(cookie, c, 2).json()["job"]
        self.w.fail |= {"az.test", "api.cloudflare.com"}
        for _ in range(6):
            get_store().update("jobs", job["id"], lambda j: {**j, "leaseUntil": 0} if j else j)
            jobs.run_job(job["id"], 120_000)
        j = get_store().get("jobs", job["id"])
        self.assertEqual(j["status"], "failed")
        self.assertIn("not counted", j["error"])
        self.assertEqual(get("/api/me", cookie=cookie).json()["usage"]["jobs"], 0)

    def test_one_worker_at_a_time_and_stalled_runs_resume(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        job = self.start_job(cookie, c, 1).json()["job"]
        st = get_store()
        st.update("jobs", job["id"], lambda j: {**j, "status": "running", "leaseUntil": time.time() * 1000 + 60_000})
        self.assertEqual(jobs.run_job(job["id"]), "busy")
        self.assertEqual(jobs.due_jobs(), [])
        st.update("jobs", job["id"], lambda j: {**j, "leaseUntil": time.time() * 1000 - 1})
        self.assertEqual(jobs.due_jobs(), [job["id"]])
        self.assertEqual(jobs.run_job(job["id"], 120_000), "done")

    def test_other_users_cannot_reach_a_run_or_its_files(self):
        a, _ = self.sign_up("a@example.com")
        b, _ = self.sign_up("b@example.com")
        c = self.make_course(a)
        job = self.start_job(a, c, 1).json()["job"]
        self.run_done(a, job["id"])
        self.assertEqual(get(f"/api/jobs/{job['id']}", cookie=b).status, 403)
        self.assertEqual(get(f"/api/jobs/{job['id']}/download?format=md", cookie=b).status, 403)
        self.assertEqual(get(f"/api/jobs/{job['id']}/download?format=exe", cookie=a).status, 400)
        self.assertEqual(get("/api/jobs/..%2F..", cookie=a).status, 404)

    def test_best_selection_compares_providers(self):
        self.w.weak_azure = True
        r = generate_best("Maths", "Algebra", "Linear equations", 2)
        self.assertEqual(len(r["all"]), 2)
        self.assertEqual(r["best"]["provider"], "cloudflare")
        self.assertGreater(r["best"]["score"], min(x["score"] for x in r["all"]))


class TestImages(ApiBase):
    def test_moves_to_next_model_when_one_fails(self):
        import base64
        n = {"v": 0}

        def t(method, url, headers, body, timeout):
            if "flux-1-schnell" in url:
                return Resp(503, b"down")
            n["v"] += 1
            return Resp(200, json.dumps({"result": {"image": base64.b64encode(PNG).decode()}}).encode(), {"content-type": "application/json"})

        net.set_transport(t)
        img = generate_image("a cell")
        self.assertIn("stable-diffusion", img["model"])
        self.assertEqual(img["mime"], "image/png")
        self.assertEqual(n["v"], 1)

    def test_503_when_every_model_fails_or_returns_junk(self):
        net.set_transport(lambda *a: Resp(200, bytes([1, 2, 3, 4]), {"content-type": "image/png"}))
        with self.assertRaises(Exception) as cm:
            generate_image("x")
        self.assertEqual(cm.exception.status, 503)

    def test_gated_by_plan(self):
        cookie, _ = self.sign_up()
        self.assertEqual(post("/api/images", {"prompt": "a cell diagram"}, cookie=cookie).status, 402)
        self.upgrade("asha@example.com")
        r = post("/api/images", {"prompt": "a cell diagram"}, cookie=cookie)
        self.assertEqual(r.status, 200)
        self.assertEqual(r.headers["content-type"], "image/png")
        self.assertEqual(r.headers["x-ai-generated"], "true")


def _event(ref, **over):
    return json.dumps({"event": "payment_link.paid", "payload": {"payment_link": {"entity": {"id": "plink_1", "reference_id": ref, "status": "paid", "amount": 29900,
                                                                                         "amount_paid": 29900, "currency": "INR", **over}}, "payment": {"entity": {"id": "pay_1"}}}})


class TestPayments(ApiBase):
    def webhook(self, raw, id="evt_1", sig=None):
        return post("/api/webhooks/razorpay", raw=raw, origin=None, headers={"x-razorpay-signature": sig or hmac_hex("whsecret", raw), "x-razorpay-event-id": id})

    def checkout(self, cookie, plan="learner", billing="month"):
        return post("/api/checkout", {"plan": plan, "billing": billing}, cookie=cookie)

    def test_lists_paid_plan_with_checkout_action(self):
        learner = next(p for p in get("/api/plans", origin=None).json()["plans"] if p["id"] == "learner")
        self.assertEqual(learner["price"], 299)
        self.assertEqual(learner["cta"]["action"], "checkout")

    def test_one_hosted_link_even_if_pressed_twice(self):
        cookie, _ = self.sign_up()
        a, b = self.checkout(cookie).json(), self.checkout(cookie).json()
        self.assertEqual(a["url"], "https://rzp.io/i/abc")
        self.assertEqual(b["orderId"], a["orderId"])
        self.assertEqual(len([c for c in self.w.calls if c["url"].endswith("/payment_links") and c["method"] == "POST"]), 1)
        sent = next(c for c in self.w.calls if c["url"].endswith("/payment_links"))
        self.assertEqual(json.loads(sent["body"])["amount"], 29900)
        self.assertEqual(json.loads(sent["body"])["currency"], "INR")
        self.assertTrue(sent["headers"]["authorization"].startswith("Basic "))

    def test_purchase_is_recorded_with_payment_id_and_plan_activates_only_after_payment(self):
        cookie, _ = self.sign_up()
        oid = self.checkout(cookie).json()["orderId"]
        self.assertEqual(get("/api/purchases", cookie=cookie).json()["purchases"], [])  # nothing is recorded or active before payment
        self.assertEqual(get("/api/me", cookie=cookie).json()["user"]["plan"], "prototype")
        self.assertEqual(self.webhook(_event(oid), "evt_p").status, 200)
        rows = get("/api/purchases", cookie=cookie).json()["purchases"]
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["paymentId"], rows[0]["item"], rows[0]["status"], rows[0]["amountPaise"]), ("pay_1", "learner", "active", 29900))
        self.assertIsNotNone(rows[0]["activeUntil"])
        self.assertEqual(get("/api/me", cookie=cookie).json()["user"]["plan"], "learner")
        self.webhook(_event(oid), "evt_p2")  # a repeat changes nothing
        self.assertEqual(len(get("/api/purchases", cookie=cookie).json()["purchases"]), 1)

    def test_beta_plus_switches_on_beta_and_the_study_pack_together_at_a_lower_price(self):
        plans = {p["id"]: p for p in get("/api/plans", origin=None).json()["plans"]}
        self.assertEqual((plans["plus"]["price"], plans["plus"]["listPrice"]), (179, 428))  # tests set Beta at 299, the pack costs 129
        self.assertGreater(plans["plus"]["saveInr"], 0)
        cookie, _ = self.sign_up()
        oid = self.checkout(cookie, plan="plus").json()["orderId"]
        self.assertEqual(self.webhook(_event(oid, amount=17900, amount_paid=17900), "evt_plus").status, 200)
        me = get("/api/me", cookie=cookie).json()
        self.assertEqual(me["user"]["plan"], "learner")
        mine = get("/api/addons", cookie=cookie).json()["mine"]
        self.assertEqual(sorted(m["sku"] for m in mine), ["chatbot", "monthly_notes"])
        self.assertEqual(get("/api/purchases", cookie=cookie).json()["purchases"][0]["item"], "plus")

    def test_refuses_signed_out_free_plan_or_payments_off(self):
        self.assertEqual(post("/api/checkout", {"plan": "learner", "billing": "month"}).status, 401)
        cookie, _ = self.sign_up()
        self.assertEqual(self.checkout(cookie, "prototype").status, 400)
        self.assertEqual(self.checkout(cookie, "classroom").status, 400)
        self.w = setup(RAZORPAY_KEY_ID="", RAZORPAY_KEY_SECRET="")
        c2, _ = self.sign_up("z@example.com")
        self.assertEqual(self.checkout(c2).status, 503)

    def test_activates_from_signed_webhook_once_even_when_replayed(self):
        cookie, _ = self.sign_up()
        order = self.checkout(cookie).json()["orderId"]
        raw = _event(order)
        self.assertEqual(self.webhook(raw).status, 200)
        u1 = self.user_doc("asha@example.com")
        self.assertEqual(effective_plan(u1), "learner")
        until = u1["planUntil"]
        self.assertGreater(until - time.time() * 1000, 29 * 86_400_000)
        self.assertEqual(self.webhook(raw).json()["result"], "duplicate")
        self.assertEqual(self.webhook(raw, "evt_2").status, 200)  # Razorpay can resend with a new event id: must not extend again
        self.assertEqual(self.user_doc("asha@example.com")["planUntil"], until)
        paid = [c for c in self.w.calls if c["url"] == "https://n8n.test/hook" and "payment.paid" in c["body"]]
        self.assertEqual(len(paid), 1)

    def test_rejects_bad_signatures_wrong_amounts_unknown_orders(self):
        cookie, _ = self.sign_up()
        order = self.checkout(cookie).json()["orderId"]
        self.assertEqual(self.webhook(_event(order), "e1", "deadbeef").status, 400)
        self.assertEqual(self.webhook(_event(order, amount_paid=100), "e2").status, 409)
        self.assertEqual(effective_plan(self.user_doc("asha@example.com")), "prototype")
        self.assertEqual(self.webhook(_event("ord_missing"), "e3").status, 404)
        self.assertEqual(self.webhook(_event(order), "e2").status, 200)  # a failed attempt must not burn the event id
        self.assertEqual(effective_plan(self.user_doc("asha@example.com")), "learner")

    def test_renewal_extends_from_current_end(self):
        cookie, _ = self.sign_up()
        a = self.checkout(cookie).json()["orderId"]
        self.webhook(_event(a), "ev_a")
        first = self.user_doc("asha@example.com")["planUntil"]
        get_store().update("orders", a, lambda o: {**o, "status": "expired", "created": 0})
        b = self.checkout(cookie, billing="year").json()["orderId"]
        self.webhook(_event(b, amount=299900, amount_paid=299900), "ev_b")
        self.assertGreater(self.user_doc("asha@example.com")["planUntil"] - first, 364 * 86_400_000)

    def _return_qs(self, order, sig=None):
        qs = {"razorpay_payment_link_id": "plink_1", "razorpay_payment_link_reference_id": order, "razorpay_payment_link_status": "paid", "razorpay_payment_id": "pay_9"}
        qs["razorpay_signature"] = sig or hmac_hex("keysecret", f"plink_1|{order}|paid|pay_9")
        return urlencode(qs)

    def test_confirms_redirect_with_razorpay_and_activates(self):
        cookie, _ = self.sign_up()
        order = self.checkout(cookie).json()["orderId"]
        r = get(f"/api/billing/return?{self._return_qs(order)}", cookie=cookie)
        self.assertEqual(r.status, 303)
        self.assertEqual(r.headers["location"], "http://localhost:3000/app/settings?paid=1")
        self.assertEqual(effective_plan(self.user_doc("asha@example.com")), "learner")
        self.assertIn("paid=invalid", get(f"/api/billing/return?{self._return_qs(order, 'nope')}", cookie=cookie).headers["location"])

    def test_ignores_unpaid_redirect(self):
        cookie, _ = self.sign_up()
        order = self.checkout(cookie).json()["orderId"]
        self.w.paid = False
        r = get(f"/api/billing/return?{self._return_qs(order)}", cookie=cookie)
        self.assertIn("paid=0", r.headers["location"])
        self.assertEqual(effective_plan(self.user_doc("asha@example.com")), "prototype")

    def test_expired_plan_falls_back_to_free(self):
        now = time.time() * 1000
        self.assertEqual(effective_plan({"plan": "learner", "planUntil": now - 1}), "prototype")
        self.assertEqual(effective_plan({"plan": "learner", "planUntil": now + 1000}), "learner")


class TestFeedbackGate(ApiBase):
    def test_blocks_next_session_until_rated_and_cannot_be_skipped(self):
        cookie, _ = self.sign_up()
        self.upgrade("asha@example.com")
        s1 = post("/api/sessions", {"courseId": "c1", "topicId": "t1"}, cookie=cookie)
        self.assertEqual(s1.status, 201)
        sid = s1.json()["session"]["id"]
        blocked = post("/api/sessions", {}, cookie=cookie)
        self.assertEqual(blocked.status, 403)
        self.assertEqual(blocked.json()["error"], "feedback_required")
        self.assertEqual(blocked.json()["sessionId"], sid)
        self.assertTrue(get("/api/me", cookie=cookie).json()["feedbackRequired"])
        self.assertEqual(post("/api/feedback", {"sessionId": sid, "rating": 9}, cookie=cookie).status, 422)
        ok = post("/api/feedback", {"sessionId": sid, "rating": 5, "comment": "Loved it", "tags": ["content", "bogus"]}, cookie=cookie)
        self.assertEqual(ok.status, 201)
        self.assertEqual(post("/api/feedback", {"sessionId": sid, "rating": 1}, cookie=cookie).status, 200)  # idempotent
        self.assertEqual(post("/api/sessions", {}, cookie=cookie).status, 201)
        self.assertEqual(len([c for c in self.w.calls if "feedback.received" in c["body"]]), 1)

    def test_whatsapp_after_low_rating_and_private_sessions(self):
        a, _ = self.sign_up("a@example.com")
        b, _ = self.sign_up("b@example.com")
        self.upgrade("a@example.com")
        sid = post("/api/sessions", {}, cookie=a).json()["session"]["id"]
        self.assertEqual(post("/api/feedback", {"sessionId": sid, "rating": 3}, cookie=b).status, 403)
        low = post("/api/feedback", {"sessionId": sid, "rating": 1, "comment": "broken"}, cookie=a).json()
        self.assertRegex(low["whatsapp"], r"^https://wa\.me/919876543210\?text=")

    def test_grace_allowance(self):
        self.w = setup(FEEDBACK_GRACE_SESSIONS="1")
        cookie, _ = self.sign_up()
        self.upgrade("asha@example.com")
        codes = [post("/api/sessions", {}, cookie=cookie).status for _ in range(3)]
        self.assertEqual(codes, [201, 201, 403])


class TestSupport(ApiBase):
    def test_ticket_whatsapp_and_n8n(self):
        cookie, _ = self.sign_up()
        r = post("/api/support", {"message": "My deck did not download"}, cookie=cookie).json()
        self.assertRegex(r["ticketId"], r"^t_")
        from urllib.parse import unquote
        self.assertIn(r["ticketId"], unquote(r["whatsapp"]))
        self.assertTrue(any("support.requested" in c["body"] for c in self.w.calls))

    def test_visitors_and_leads(self):
        self.assertEqual(post("/api/support", {"message": "Pricing question please", "contact": "x@y.com"}).status, 201)
        self.assertEqual(post("/api/leads", {"name": "Ravi", "email": "ravi@example.com", "role": "Teacher", "message": "Demo please"}).status, 201)
        self.assertEqual(post("/api/leads", {"name": "", "email": "bad"}).status, 422)

    def test_ops_whatsapp_alert(self):
        self.w = setup(WHATSAPP_TOKEN="t", WHATSAPP_PHONE_ID="123", WHATSAPP_OPS_TO="911111111111")
        post("/api/support", {"message": "Need help with login"})
        c = next(c for c in self.w.calls if "graph.facebook.com/v21.0/123/messages" in c["url"])
        self.assertEqual(json.loads(c["body"])["to"], "911111111111")


class TestN8n(ApiBase):
    def n8n(self, b, headers=None):
        return post("/api/automation/n8n", b, origin=None, headers={"authorization": "Bearer n8nsecret"} if headers is None else headers)

    def test_signs_outbound_events(self):
        self.sign_up()
        c = next(c for c in self.w.calls if c["url"] == "https://n8n.test/hook")
        ts = c["headers"]["x-wlxc-timestamp"]
        self.assertEqual(c["headers"]["x-wlxc-signature"], integ.sign("n8nsecret", ts, c["body"]))
        self.assertEqual(json.loads(c["body"])["type"], "user.registered")

    def test_retries_failed_deliveries_from_outbox(self):
        self.w.fail.add("n8n.test")
        self.sign_up()
        st = get_store()
        pending = st.find("outbox", where={"status": "pending"})
        self.assertEqual(len(pending), 1)
        self.w.fail.discard("n8n.test")
        st.update("outbox", pending[0]["id"], lambda c: {**c, "next": 0})
        self.assertEqual(integ.flush_outbox()["sent"], 1)
        self.assertEqual(len(st.find("outbox", where={"status": "pending"})), 0)

    def test_outage_never_breaks_signup(self):
        self.w.fail.add("n8n.test")
        self.assertEqual(self.sign_up()[1].status, 201)

    def test_authenticates_inbound(self):
        self.assertEqual(self.n8n({"action": "ping"}, {}).status, 401)
        self.assertEqual(self.n8n({"action": "ping"}, {"authorization": "Bearer wrong"}).status, 401)
        self.assertEqual(self.n8n({"action": "ping"}).status, 200)
        raw = json.dumps({"action": "ping"})
        ts = str(int(time.time()))
        self.assertEqual(post("/api/automation/n8n", raw=raw, origin=None, headers={"x-wlxc-timestamp": ts, "x-wlxc-signature": integ.sign("n8nsecret", ts, raw)}).status, 200)
        stale = str(int(time.time()) - 3600)
        self.assertEqual(post("/api/automation/n8n", raw=raw, origin=None, headers={"x-wlxc-timestamp": stale, "x-wlxc-signature": integ.sign("n8nsecret", stale, raw)}).status, 401)

    def test_grant_plan_and_create_job_by_email(self):
        cookie, _ = self.sign_up("kid@example.com")
        c = self.make_course(cookie)
        self.assertEqual(self.n8n({"action": "grant_plan", "email": "kid@example.com", "plan": "learner", "days": 10}).status, 200)
        self.assertEqual(effective_plan(self.user_doc("kid@example.com")), "learner")
        self.assertEqual(self.n8n({"action": "create_job", "email": "kid@example.com", "courseId": c["id"], "topicIds": c["course"]["order"][:1]}).status, 202)
        self.assertEqual(self.n8n({"action": "grant_plan", "email": "ghost@example.com", "plan": "learner", "days": 1}).status, 404)


class TestCronAndHealth(ApiBase):
    def test_rejects_callers_without_secret(self):
        self.assertEqual(get("/api/internal/cron", origin=None).status, 401)
        self.assertEqual(get("/api/internal/cron", origin=None, headers={"authorization": "Bearer nope"}).status, 401)

    def test_resumes_runs_and_reminds_once(self):
        cookie, _ = self.sign_up()
        c = self.make_course(cookie)
        job = self.start_job(cookie, c, 1).json()["job"]
        st = get_store()
        st.update("jobs", job["id"], lambda j: {**j, "status": "queued", "leaseUntil": 0})
        u = self.user_doc("asha@example.com")
        st.update("users", u["id"], lambda x: {**x, "plan": "learner", "planUntil": time.time() * 1000 + 2 * 86_400_000})
        r1 = get("/api/internal/cron", origin=None, headers=AUTH_CRON).json()
        self.assertEqual(r1["jobs"][job["id"]], "done")
        self.assertEqual(r1["reminded"], 1)
        self.assertEqual(get("/api/internal/cron", origin=None, headers=AUTH_CRON).json()["reminded"], 0)

    def test_health_minimal_public_detailed_with_secret(self):
        self.assertEqual(get("/api/health", origin=None).json(), {"ok": True})
        j = get("/api/health", origin=None, headers=AUTH_CRON).json()
        self.assertTrue(j["ok"])
        self.assertTrue(j["integrations"]["azure"])
        self.assertNotIn("keysecret", json.dumps(j))


if __name__ == "__main__":
    unittest.main()
