#!/usr/bin/env python3
"""Builds the WorkLab XC guide as a PDF and a Word file from one source:  python scripts/make_docs.py
Needs reportlab and python-docx (both are in backend/requirements.txt). Output: docs/WorkLab-XC-Guide.pdf and .docx.
Edit CONTENT below and run it again after the product changes."""
from __future__ import annotations

import re
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "docs"

# ------------------------------------------------------------------------------------------------------------------ content
# ("h1", text) ("h2", text) ("h3", text) ("p", text) ("b", [items]) ("code", text) ("note", text) ("t", rows, widths_mm)
# Inline: **bold** and `code`.
C: list[tuple] = []


def H1(t): C.append(("h1", t))
def H2(t): C.append(("h2", t))
def H3(t): C.append(("h3", t))
def P(t): C.append(("p", t))
def B(*items): C.append(("b", list(items)))
def CODE(t): C.append(("code", t.strip("\n")))
def NOTE(t): C.append(("note", t))
def T(rows, widths): C.append(("t", rows, widths))
def PB(): C.append(("pb",))


# ------------------------------ cover and contents
C.append(("title", "WorkLab XC", "Complete guide: how to use it, how to set it up, and how to publish it online"))
P("Syllabus to AI lessons to a 3D lecture hall. This guide has three parts: **using the platform** (for learners), **setting it up and configuring it** (for the owner), and **publishing it online**, followed by how to customise it and how to fix common problems.")
T([["Part", "For", "What it covers"],
   ["1. Using WorkLab XC", "Learners", "Account, uploading a syllabus, lessons, listening, podcasts, practice rounds, the 3D classroom, exams, plans, support"],
   ["2. Setting it up", "Owner", "Requirements, install, run, every setting in .env, keys for AI, email, payments, Google sign-in, WhatsApp support, ads"],
   ["3. Publishing online", "Owner", "Docker, Render, Fly.io, Vercel, a plain server, domain and HTTPS, go-live checklist, backups and updates"],
   ["4. Customising", "Owner", "Plans and prices, the 3D hall and presenter, small-screen design, wording and AI prompts"],
   ["5. Help", "Everyone", "Troubleshooting, command cheat sheet, honest limits"]], [38, 24, 108])
PB()

# ------------------------------ PART 1
H1("Part 1. Using WorkLab XC")
H2("1.1 What it does")
B("You give it a **syllabus**. It builds a **course map** of units and topics, and writes a **lesson** for each topic with AI.",
  "You study each topic in a **lesson view** (slides, narration you can listen to) or in the **3D classroom**.",
  "You then **practise**: a practice round of quick cards, a game, worked problems, a practice exam, and questions to the AI tutor.",
  "Everything you do updates a **mastery estimate** for each topic, so the app can suggest what to study next. It is an estimate to guide you, not a grade.")
NOTE("AI can make mistakes. Check anything important against your textbook, and use the feedback and report buttons when something looks wrong.")

H2("1.2 Create an account and sign in")
B("Open the site and choose **Create account** (or **Get started**). You can sign up with email and password, or with **Continue with Google** when it is switched on.",
  "Fill in every detail: name, email, password, age, what you are studying at, syllabus type, class or year, and school or institute. These help match lessons to your level and syllabus, and you can change them later in **Profile**.",
  "Tick the box for the Terms, Privacy policy and AI use notice. Confirm your email using the link we send you (check spam if it does not arrive).",
  "You stay signed in for about 30 days on the same device. When you open the home page while signed in, you go straight to your workspace. Use **Sign Out** in the account menu on shared computers.",
  "Forgot your password? Use **Forgot password?** on the sign-in page.")

H2("1.3 Finding your way around")
T([["Where", "What you find"],
   ["Sidebar", "Dashboard, Upload syllabus, Learn, Saved, Courses, Knowledge graph, Practice, AI tutor, Progress. At the bottom: your streak and points, an Upgrade pill (Free accounts) and your account card."],
   ["Account card (bottom of the sidebar)", "Profile, Settings, Help centre, Support, Download app (installs WorkLab XC like an app) and Sign Out."],
   ["Top bar", "Search (press /), theme button (dark or light), notifications bell, and your photo."],
   ["Light and dark", "Use the theme button. Light mode uses warm cream panels; dark mode is the original dark look."]], [46, 124])

H2("1.4 Upload a syllabus")
B("Go to **Upload syllabus**. Drop a file, choose a file, or paste the text into the box, then press **Build course**.",
  "**Files:** text and Markdown on every plan; PDF and Word on Beta and Beta Plus. Maximum 5 MB.",
  "**Best layout:** each unit as a heading and each topic on its own line, for example `## Unit 1: Motion` then `- Velocity`. Page numbers like (p. 3) are kept.",
  "A progress card shows the upload, reading and building steps. When it finishes you land on your **course map**, where you can check the units and topics and confirm or remove prerequisite links that were guessed.",
  "**Free** keeps one subject at a time. **Beta** allows many subjects.")

H2("1.5 The Learn page: your learning paths")
B("**Learn** shows every subject as its own board, side by side. Each board has a header with your progress, then each unit with its topics down a timeline.",
  "Topics unlock **one at a time**. Finish a topic (study it and reach 60% on its checks) to open the next. The current topic is marked **Start** and **Up next**.",
  "Press **Podcast** on a board to make an audio podcast of any open topic (section 1.8).",
  "Tap a topic to see its estimate, then **Start**, **Review** or **Practise**.")

H2("1.6 A learning session")
P("Press **Start** on a topic. The first time, the lesson is written by AI, which can take up to a minute. After that it opens instantly, and lessons already written for other learners are reused. The session has five tabs:")
T([["Tab", "What it does"],
   ["Lesson", "Slides with key points and a spoken script. The **Listen** player at the top reads the lesson aloud: play and pause, previous and next slide, speed (1x, 1.25x, 1.5x, 0.75x) and a progress bar. Three check questions at the side update your estimate."],
   ["Practice round", "Quick cards: multiple choice, fill in the blank and true or false (section 1.7). Beta and Beta Plus."],
   ["Game", "True or false, put-in-order and matching rounds, checked on the server. Beta and Beta Plus."],
   ["Practice problems", "Worked examples and numeric problems with hints first, then the full solution when you ask. Beta and Beta Plus."],
   ["Ask the class", "Questions to the AI tutor, answered from this lesson with the slide it used. 2 free questions a month; unlimited with the tutor add-on or Beta Plus."]], [34, 136])

H2("1.7 Practice rounds: quizzes, fill in the blank, timed")
B("Open the **Practice round** tab, or press **Practice round** when a classroom lesson finishes.",
  "A round has 6 to 10 cards. Pick an answer and press **Check**. You see at once whether you were right and **why**, then press **Continue**.",
  "Fill-in-the-blank cards show the sentence with a box that fills with the word you choose.",
  "Tick **Timed round** for 20 seconds a card with bonus points for speed; a card left too long counts as wrong.",
  "You earn points for correct answers, plus a bonus for answering several in a row. At the end you see your score, points, time, and the cards worth another look, with their explanations.",
  "Every answer updates the topic's estimate. Press **Practise again** to replay the same cards or **New round** for a fresh set.")

H2("1.8 Podcasts")
B("On the **Learn** page press **Podcast** on a subject. Choose a style under **Customize**: **Deep dive** (about six minutes) or **Quick recap** (about two minutes).",
  "Press **Generate** beside an open topic. We write a short two-host conversation (Asha and Ravi) about it and your browser reads it aloud with two voices.",
  "Controls: play and pause, previous and next line, speed, a progress bar and a transcript that highlights the line being read.",
  "**Allowance:** every account can make **2 podcasts a month for free**. A topic that has already been made, by you or by anyone, opens instantly and costs nothing. The **Podcast pass** add-on (or **Beta Plus**) removes the limit.",
  "Voices come from your browser, so quality depends on your device. If your browser cannot read aloud, the conversation is shown as text.")

H2("1.9 The 3D classroom")
B("Press **Enter classroom** on a lesson. It is a lecture hall with a **presentation screen** that runs your slides, a **left board** that explains in words, a **right board** that shows the live question or formula, and a presenter who moves while the lesson plays.",
  "**Controls:** Space plays or pauses; keys 1, 2 and 3 move the camera to the board, model or overview; drag to look around; the **+**, **-** and **fit** buttons on the stage zoom safely to the board (you cannot zoom out of the room or lose the screen). Pinch also works on a phone.",
  "**On a phone:** in portrait the classroom is a video-style player at the top with the study tabs below it. Turn the phone sideways for a bigger stage with a slim control bar.",
  "Captions and spoken narration can be switched on or off with the stage buttons (keys C and M). Playback speed is on the bar.",
  "When the lesson ends, choose **Practice round** (a quick quiz without leaving the room), **Replay**, the next topic, or **Play revision game**.",
  "**Allowance:** Free includes one classroom session a month (reopening the same topic is free); Beta and Beta Plus are unlimited. If the classroom looks black, see the Help section in Part 5.")

H2("1.10 Knowledge graph")
B("Open **Knowledge graph**. The default **Tree** view shows subjects, units, topics and the parts inside each topic. **Click any item to open it in place**; use **Open all** and **Close all** to expand or collapse everything. Each item has a progress ring and a status.",
  "Click a topic to see its summary, links to study it, practise it or open it in the classroom, and the prerequisite links we guessed from the order of your syllabus (confirm or remove them).",
  "The **Map** view shows the same information as circles and arrows: all subjects, then the units of one subject, then the topics of a unit and how they depend on each other.")

H2("1.11 Practice exams and handwritten answers")
B("Beta and Beta Plus include **practice exams**. Choose a course, topics, number of questions, time limit and question types.",
  "Sit the paper in a **timed exam window**. Answer by typing, or upload a clear photo of your handwritten answer sheet to be read and marked by AI.",
  "Answers and explanations are shown only after marking. AI marks are a guide for your practice, not an official grade. If handwriting is hard to read you are told, and a clearer photo helps.")

H2("1.12 Notes by marks and full answers")
B("In **Learn**, open **Unit tools** to make **notes by marks** for a unit: give the marks each topic carries and get notes sized to match, with a last-day revision checklist. Each account includes **2 a month for free**; the **Unlimited notes** add-on removes the limit.",
  "**Full answers** writes model exam answers for a whole unit. Each version is saved and can be downloaded.")

H2("1.13 AI tutor, Progress and Saved")
B("**AI tutor** answers questions from a topic's lesson. 2 free questions a month; unlimited with the add-on or Beta Plus.",
  "**Progress** shows your activity graph (days you studied, like a contribution graph), your streak, study minutes, and your plan's usage for today and this month.",
  "**Saved** lists everything you studied by subject, unit and topic, with your questions and view-only notes.")

H2("1.14 Plans, add-ons and buying")
T([["", "Free", "Beta", "Beta Plus"],
   ["Learning sessions", "Unlimited", "Unlimited", "Unlimited"],
   ["Subjects", "1 at a time", "Unlimited", "Unlimited"],
   ["Ads", "Yes, including video", "None", "None"],
   ["3D classroom", "1 session a month", "Unlimited", "Unlimited"],
   ["Games, quizzes, practice rounds, exams", "No", "Yes", "Yes"],
   ["PDF and Word syllabi", "No", "Yes", "Yes"],
   ["Each month free", "2 tutor questions, 2 unit notes, 2 podcasts", "same", "Unlimited tutor, notes and podcasts"]], [52, 40, 38, 40])
B("**How to buy:** on **Pricing** choose Monthly or Yearly (yearly shows the saving), optionally tick **Renew automatically**, and press **Buy**. If you are signed out you are asked to create an account or sign in first and then taken straight on to payment. You pay on a secure **Razorpay** page with UPI, cards or netbanking; we never see your card details.",
  "**Add-ons** (Pricing, \"Add what you need\"): Unit notes by marks, Unlimited notes (30 days), AI tutor chat (30 days), Podcast pass (30 days) and the Study pack (tutor and unlimited notes together, cheaper than buying them apart).",
  "**After paying:** your plan switches on within a minute. **Settings, Purchases** lists every payment with its **payment id** and what it activated. A receipt is emailed.",
  "**Cancel auto-renewal** any time in Settings. When a plan ends your account returns to Free and nothing is deleted.")

H2("1.15 Support")
B("Open **Support** from the account menu. Write a message and we reply in the same chat. The team sees only a short code, your plan and your message, never your name, email or phone, and you never see theirs.",
  "To talk something through, **book a session**: pick a time that suits you. We confirm by email, remind you an hour before, and the session happens in the chat. Free accounts can book 1 a month and Beta accounts 4.",
  "The **Help centre** (/help) answers about 60 common questions with search.")

H2("1.16 Settings, privacy and your data")
B("**Settings** has theme, text size, motion, 3D quality (Low, Medium, High), captions, spoken narration, playback speed, notifications, product emails, ad choice, usage, purchases and account actions.",
  "**Privacy:** lessons you generate can be shared with other learners studying the same topic name. Do not put personal details in a topic title. Download everything we hold about you, or delete your account, from Settings.",
  "Product emails can be switched off in Settings or with the one-click unsubscribe link. Receipts, security and plan-ending notices are always sent.")
PB()

# ------------------------------ PART 2
H1("Part 2. Setting it up and configuring it")
H2("2.1 What you need")
T([["Item", "Requirement"],
   ["Computer", "Windows 10/11, macOS 12+, or Linux. About 2 GB free disk and 4 GB RAM."],
   ["Python", "3.10 or newer (3.12 recommended). Installed for you if missing."],
   ["Node.js", "20.9 or newer (22 LTS recommended). Installed for you if missing."],
   ["Internet", "To install packages, and at run time for AI, email, payments and Google."],
   ["Browser for users", "Current Chrome, Edge, Firefox or Safari. WebGL is needed for the 3D classroom (a 2D lesson view is the fallback)."],
   ["Ports", "3000 (web) and 8000 (API) free."],
   ["Accounts you will want", "An AI key (Gemini is free to start), a mail provider (Resend), Razorpay (for money), optionally Google Cloud (sign-in), Meta WhatsApp Cloud API (support), Google AdSense (ads)."]], [38, 132])

H2("2.2 Install")
P("Open the project folder and use whichever is easiest. They all run the same installer: it checks your computer, installs missing tools and every package, creates the **.env** settings file with fresh random secrets, and runs the backend self-test. It never overwrites an existing .env.")
T([["Computer", "Install (once)", "Run"],
   ["Windows", "Double-click INSTALL-Windows.bat", "Double-click START-Windows.bat"],
   ["Mac", "Double-click INSTALL-Mac.command (first time: right-click, Open)", "Double-click START-Mac.command"],
   ["Linux", "bash INSTALL-Linux.sh", "bash START-Linux.sh"]], [26, 92, 52])
CODE("""python scripts/install.py --admin-email you@example.com   # install everything   (npm run setup)
python scripts/install.py --doctor                        # only check this computer (npm run doctor)
python scripts/run.py --open                              # start API + web and open the browser (npm run app)
python scripts/run.py --prod                              # production build and run (npm run app:prod)
python scripts/run.py --api-only                          # API only, port 8000
python scripts/run.py --check                             # start, run the staging checks, stop""")
P("The first start is slower because the web app compiles. Press Ctrl-C to stop. Open http://localhost:3000.")

H2("2.3 Become the administrator")
B("Put your email in **ADMIN_EMAILS** in .env (the installer's `--admin-email` does this), restart, and sign up or sign in with that email. The **Admin** page appears in your account menu.",
  "To create a ready-made account (for yourself, a partner or a tester) with the paid plan switched on and no payment:")
CODE('python scripts/with_env.py -m wlxc.cli_account someone@example.com "a-strong-password" --name "Their Name" --plus [--admin]')

H2("2.4 The .env file: every setting")
P("All settings are in **.env** in the project folder. One `KEY=value` per line, no quotes. **Restart after editing.** Settings that start with **NEXT_PUBLIC_** are fixed when the web app is built, so rebuild after changing them. **.env.example** lists everything with comments. Never share .env.")
T([["Setting", "Meaning"],
   ["APP_URL", "Public address of the site. http://localhost:3000 locally, https://yourdomain.com online."],
   ["API_ORIGIN", "Where the web app forwards /api/* to (http://127.0.0.1:8000 locally)."],
   ["SESSION_SECRET, CRON_SECRET", "Long random strings (generated for you). Changing SESSION_SECRET signs everyone out."],
   ["WLXC_ENV", "development or production. Production turns on strict defaults."],
   ["SQLITE_PATH / DATABASE_URL", "SQLite file (one server) or a PostgreSQL connection string (production)."],
   ["WORKER_MODE", "inline (jobs run inside the API, simplest) or external (a separate worker process)."],
   ["ADMIN_EMAILS", "Comma-separated emails that become administrators."],
   ["REQUIRE_SIGNUP_PROFILE", "1 (default) makes age, education, syllabus type, class and institute mandatory at sign-up."],
   ["REQUIRE_CONSENT", "1 (default) requires the terms and privacy box."],
   ["FAIR_GLOBAL_HEAVY", "Most heavy AI requests allowed at once across the server (default 60)."],
   ["MAX_UPLOAD_BYTES", "Largest syllabus or sheet upload."]], [52, 118])

H2("2.5 AI keys (needed for lessons, practice rounds, podcasts, exams and the tutor)")
P("Without a key, lessons show \"no AI provider is set up\" instead of fake content. Add at least one; list several in **LLM_ORDER** as automatic fallbacks.")
T([["Provider", "Where to get a key", ".env"],
   ["Google Gemini (free tier)", "aistudio.google.com/apikey", "GEMINI_API_KEY="],
   ["Groq, Mistral, OpenRouter", "console.groq.com, console.mistral.ai, openrouter.ai/keys", "GROQ_API_KEY, MISTRAL_API_KEY, OPENROUTER_API_KEY"],
   ["OpenAI or Azure OpenAI", "platform.openai.com, Azure portal", "OPENAI_API_KEY; AZURE_OPENAI_ENDPOINT, _API_KEY, _DEPLOYMENT"],
   ["Cloudflare Workers AI", "dash.cloudflare.com", "CF_ACCOUNT_ID, CF_API_TOKEN"]], [46, 66, 58])
B("Handwritten answer sheets need a provider that accepts images (Gemini, OpenAI or Azure).",
  "If learners upload private notes, put `azure` or `openai` first in LLM_ORDER: some free tiers may train on prompts.",
  "After adding a key, open a topic in the app to confirm a lesson is written.")

H2("2.6 Email from your own domain")
B("Create an account at resend.com, add your domain (for example worklabxc.com), add the DNS records it shows at your registrar, and wait for **Verified**. Create an API key and set `RESEND_API_KEY=re_...`.",
  "Set the sender addresses on that domain: `MAIL_NOREPLY` (system mail), `MAIL_UPDATES` (announcements, default `WorkLab XC <hello@worklabxc.com>`), `SUPPORT_EMAIL`. A backup provider can be added with `BREVO_API_KEY`, `MAILGUN_*` or `SMTP_URL`.",
  "Header pictures for emails are in `public/mail/`. Redraw them with `python scripts/make_mail_images.py` (needs `pip install pillow`).",
  "Without a mail key, emails are printed in the API terminal in development.")

H2("2.7 Payments with Razorpay")
B("Create a Razorpay account, stay in **Test mode** first, and generate API keys: `RAZORPAY_KEY_ID`, `RAZORPAY_KEY_SECRET`.",
  "Create a webhook in the Razorpay dashboard: URL `https://yourdomain.com/api/webhooks/razorpay`, a secret of your choice (put it in `RAZORPAY_WEBHOOK_SECRET`), and enable the payment link and subscription events.",
  "Prices are settings, not code: `PRICE_LEARNER_MONTH_INR` / `_YEAR_INR` (Beta, default 99 / 799), `PRICE_PLUS_MONTH_INR` / `_YEAR_INR` (Beta Plus, 179 / 1499), `PRICE_ADDON_UNIT_NOTES_INR` (29), `_MONTHLY_NOTES_` (99), `_CHATBOT_MONTH_` (59), `_PODCAST_MONTH_` (49), `_PACK_MONTH_` (129). Discounts and \"you save\" amounts are calculated automatically.",
  "Test one full purchase in test mode (buy, webhook, plan active, receipt email, entry in Settings, Purchases), then switch to live keys.",
  "Until keys are set, the Buy buttons show \"Payments are not switched on yet\".")

H2("2.8 Sign in with Google")
B("In Google Cloud Console open **APIs and services, Credentials, Create credentials, OAuth client ID**, type **Web application**.",
  "Under **Authorised JavaScript origins** add your site address (and `http://localhost:3000` for testing).",
  "Put the client ID in .env twice: `GOOGLE_CLIENT_ID=` and `NEXT_PUBLIC_GOOGLE_CLIENT_ID=`. Rebuild the web app, because the second is read at build time. \"Continue with Google\" then appears on the sign-in and create-account pages.",
  "A new Google user is still asked for the same few details as email sign-up. Google-confirmed emails are marked verified.")

H2("2.9 Support chat over WhatsApp (numbers stay private)")
P("Learners chat inside the app. You reply on the **Admin page, Support desk**, or from your own WhatsApp. Neither side ever sees the other's number or email: you see only a code like C-AB12CD, the plan and the message.")
B("Set up the **Meta WhatsApp Cloud API**: `WHATSAPP_TOKEN`, `WHATSAPP_PHONE_ID`, `WHATSAPP_OPS_TO` (your own number, digits only), `WHATSAPP_APP_SECRET` and `WHATSAPP_VERIFY_TOKEN`.",
  "In the Meta dashboard set the webhook to `https://yourdomain.com/api/webhooks/whatsapp` with the same verify token and subscribe to **messages**. Only messages from WHATSAPP_OPS_TO are accepted, and every call is checked with the app secret.",
  "**Reply on WhatsApp:** send `C-AB12CD your answer` (or `C-AB12CD close`). **Session requests:** `S-FF31A0 confirm 2026-10-20 16:00` (India time), `S-FF31A0 decline reason`, `S-FF31A0 done`.",
  "`WHATSAPP_PUBLIC_LINK=0` (default) keeps your number off the site. Without WhatsApp configured, the Admin page still works on its own.")

H2("2.10 Ads (Google AdSense)")
B("Ads show only on **Free** accounts, only after the visitor picks a cookie and ads choice, and never on Beta.",
  "Set `NEXT_PUBLIC_ADSENSE_CLIENT=ca-pub-...` and the slot ids `NEXT_PUBLIC_ADSENSE_SLOT_DASHBOARD`, `_LESSON`, `_PANEL`. Use `NEXT_PUBLIC_ADSENSE_TEST=1` for test ads. Rebuild after changing them.",
  "Add your domain in the AdSense console and wait for approval before expecting ads. Video ads are served by AdSense in the same slots when enabled in your AdSense account.")

H2("2.11 Free allowances and limits")
T([["Setting", "Default", "Meaning"],
   ["FREE_CHAT_PER_MONTH", "2", "Free tutor questions a month."],
   ["FREE_NOTES_PER_MONTH", "2", "Free unit notes a month."],
   ["FREE_PODCAST_PER_MONTH", "2", "Free podcasts a month."],
   ["CHAT_MESSAGES_PER_DAY", "40", "Daily cap on tutor questions for add-on owners."],
   ["PODCASTS_PER_DAY", "30", "Daily cap on podcasts for Podcast pass owners."],
   ["LIMITS in backend/wlxc/plans.py", "-", "Per-plan limits: classroom sessions a month, subjects, games and quizzes, exams a day, syllabus size, PDF and Word, simultaneous heavy requests."]], [58, 20, 92])
P("Every limit is enforced on the server, so it cannot be bypassed from the browser. Learners see their usage under **Progress** and **Settings**.")

H2("2.12 The Admin page")
B("**Overview:** people, runs, ratings, revenue, open tickets.",
  "**Support desk:** chats by code, reply and close, session requests (confirm, decline, done), average first-reply time.",
  "**Email learners about an update:** write a subject, headline and message, choose a header picture and a button that points to your own site, choose Everyone, Free or Beta, send a test to yourself first, then send. See sent, opened, clicked and unsubscribed. At most one product email every 3 days per person.",
  "**People:** search accounts, see plan and activity, grant add-ons, block or unblock.",
  "**Purchases:** `/api/admin/purchases` lists payments with payment ids; search by payment id, order id or user id.")
P("Automatic emails (on a schedule): a renewal reminder a week before a Beta plan ends, a short note when it ends, and a gentle \"go further with Beta\" note to active Free learners at most every two weeks.")

H2("2.13 Database and backups")
B("**SQLite** (default) is one file at `SQLITE_PATH`. Fine for one server with a persistent disk. Use `WEB_CONCURRENCY=1`.",
  "**PostgreSQL** (`DATABASE_URL=postgres://...`) is recommended for production and for more than one server. Tables are created automatically (`DB_AUTO_MIGRATE`). Test it on staging first.",
  "Every record carries its owner's account id and every route checks it, so learners cannot read each other's data. Lessons and games are written once and shared by topic.",
  "**Backups:** set `BACKUP_DAILY=1`, `BACKUP_DIR`, `BACKUP_KEEP` and `BACKUP_PASSPHRASE` for encrypted daily backups, or back up the SQLite file or Postgres database yourself. Test a restore before launch.")
PB()

# ------------------------------ PART 3
H1("Part 3. Publishing it online")
H2("3.1 Before you start: the go-live checklist")
B("Buy a domain and point it at your host.",
  "Set production values: `APP_URL=https://yourdomain.com`, `WLXC_ENV=production`, long random `SESSION_SECRET` and `CRON_SECRET`, `DATABASE_URL` (or `SQLITE_PATH`), `API_ORIGIN`, `ADMIN_EMAILS`, `TRUST_PROXY=1` behind a proxy.",
  "Add at least one AI key, a verified mail domain, and for money the Razorpay live keys plus the webhook URL. Optional: Google sign-in, WhatsApp, AdSense.",
  "Use HTTPS everywhere (all options below provide it).",
  "Run the staging check (3.8) and one test payment. Read the Terms, Privacy, Refunds and Cookies pages and have them reviewed for your country: the text shipped is a default, not legal advice.")
NOTE("The platform's build and automated tests pass here. It has not been tested against live PostgreSQL, live Razorpay, live Google sign-in, live WhatsApp or live email delivery from this machine: those need your keys, so test each on staging before launch.")

H2("3.2 Option A: any server with Docker Compose (simplest all-in-one)")
CODE("""git clone <your repo> worklab-xc && cd worklab-xc
cp .env.example .env          # edit it: production values from 3.1
docker compose up --build -d  # Postgres, API, worker and web; the site is on port 3000""")
P("Put a reverse proxy with automatic HTTPS in front. With Caddy, the whole config is:")
CODE("""yourdomain.com {
    reverse_proxy 127.0.0.1:3000
}""")
P("The `NEXT_PUBLIC_*` values (API base, currency, Google client ID, AdSense) are build arguments with safe defaults, and `.dockerignore` keeps `.env` out of the image.")

H2("3.3 Option B: a plain server without Docker")
CODE("""sudo apt update && sudo apt install -y python3 python3-venv nodejs npm caddy
bash install.sh you@example.com
python3 scripts/run.py --prod      # run it as a service (systemd, pm2 or supervisor) so it restarts""")
P("Use `SQLITE_PATH` on a persistent disk, `WEB_CONCURRENCY=1`, `WORKER_MODE=inline`, and back the file up nightly.")

H2("3.4 Option C: Render (managed)")
P("New, Blueprint, pick the repository (`render.yaml` defines the API, a worker, the web app and Postgres). Fill the keys marked `sync: false` in the dashboard, and set `API_ORIGIN` on the web service to the API service's address.")

H2("3.5 Option D: Fly.io for the API, Vercel or Docker for the web")
CODE("""cd backend && fly launch --copy-config --no-deploy
fly secrets set SESSION_SECRET=... CRON_SECRET=... GEMINI_API_KEY=...
fly deploy
# web: Vercel (set API_ORIGIN=https://<api>.fly.dev) or
docker build --build-arg API_ORIGIN=https://<api>.fly.dev -t worklab-web .""")
P("The API must be a long-running service, not serverless. Use `WORKER_MODE=inline`, or call `/api/internal/cron` every 5 minutes (`.github/workflows/tick.yml`).")

H2("3.6 Domain, HTTPS and mail DNS")
B("Point your domain's A/AAAA record (or CNAME) at the host. Wait for the certificate to be issued.",
  "Add the mail provider's DNS records (SPF, DKIM) so email does not land in spam, and keep `MAIL_FROM`/`MAIL_NOREPLY` on that domain.",
  "Set the Razorpay and WhatsApp webhook URLs to the live domain, and add the live domain to Google's authorised origins and AdSense.")

H2("3.7 Keep the schedule running")
P("A small scheduler tick resumes stalled jobs, sends reminders and cleans old records. With `WORKER_MODE=inline` the API does it itself. Otherwise call `https://yourdomain.com/api/internal/cron` every 5 minutes with your `CRON_SECRET` (the GitHub workflow `tick.yml` does this).")

H2("3.8 Test before you tell anyone")
CODE("""python backend/scripts/staging_check.py https://staging.yourdomain.com --cron-secret "$CRON_SECRET"
# add --api https://api-host if the API is on another address, --api-only to skip page checks""")
B("It signs up and deletes a throw-away account and prints PASS, FAIL or SKIP for health, headers, CSRF, auth, upload, export, deletion, the scheduler and which integrations are on. Fix every FAIL.",
  "By hand: one test payment end to end; the verification and reset emails arrive; a lesson is written; the 3D classroom shows (not black) on a real laptop and a real phone; the main screens at 320, 375, 768, 1024 and 1440 px; keyboard-only use; the no-WebGL fallback.",
  "Restore a backup into a scratch database.")

H2("3.9 Running it day to day")
B("**Updating:** `git pull`, then `docker compose up --build -d` (or on a plain server: `pip install -r backend/requirements.txt`, `npm install`, `npm run build`, restart). Take a backup first.",
  "**Health:** `/api/health` answers `{\"ok\": true}`. Set `ALERT_WEBHOOK_URL` for alerts when AI providers fail or the site is slow.",
  "**Scaling:** move to PostgreSQL, run more API workers, and note that fair-use counters are per process. Heavy AI requests are limited per account and across the server.",
  "**Money:** reconcile Razorpay against **Settings, Purchases** and `/api/admin/purchases`. Every payment has a payment id.")
PB()

# ------------------------------ PART 4
H1("Part 4. Customising")
H2("4.1 Plans, prices, wording")
T([["I want to change...", "Edit"],
   ["Prices and discounts", ".env: the PRICE_* lines. \"You save\" amounts recalculate by themselves."],
   ["Limits and what each plan includes", "backend/wlxc/plans.py (LIMITS, DEFAULT_PLANS, DEFAULT_COMPARE) and the offline copy in src/lib/plans.ts. Keep both in step."],
   ["Free monthly allowances", ".env: FREE_CHAT_PER_MONTH, FREE_NOTES_PER_MONTH, FREE_PODCAST_PER_MONTH."],
   ["Page wording", "src/components/site/pages and src/components/site/landing; legal text in Legal.tsx; help answers in src/lib/faq.ts."],
   ["Colours and sizes", "Tokens at the top of src/styles/10-tokens.css, then src/styles/90-responsive.css (loaded last)."],
   ["What the AI writes (tone, length)", "The prompts in backend/wlxc/lesson.py, learn.py (practice rounds, podcasts, games) and exams.py. Restart after changing."],
   ["The AI provider or order", ".env: LLM_ORDER and the provider keys."],
   ["Add an add-on", "backend/wlxc/addons.py (SKUS, SKU_INFO, grant), a PRICE_ADDON_* setting in config.py, then it appears on Pricing and Learn automatically."]], [60, 110])

H2("4.2 The 3D engine, the hall and the presenter")
P("The classroom is drawn with **Babylon.js** (the default). The older **three.js** classroom is kept as an automatic fallback if Babylon cannot start on a device, and learners can pick either under **Settings, 3D engine**. Babylon is loaded only when a classroom opens, so it does not slow the rest of the site.")
T([["Part", "File"],
   ["Babylon scene (hall, screen, boards, presenter, models, camera, zoom, lights, frame governor)", "src/components/classroom/babylon/BabylonClassroom.tsx"],
   ["three.js fallback scene", "src/components/classroom/scene/*"],
   ["Shared measurements in metres (stage, screen, boards, presenter)", "src/components/classroom/scene/layout.ts"],
   ["Camera presets", "src/lib/classroom/timeline.ts (CAM_PRESETS)"],
   ["Hall model (lighter copy used by Babylon / original)", "public/models/hall/stage-lite.glb / stage.glb"],
   ["Presenter (GLB used by Babylon / FBX used by the fallback)", "public/models/presenter.glb / presenter.fbx"]], [100, 70])
B("**New hall:** export a `.glb`, put it in `public/models/hall/`, and point `HALL_SRC` in BabylonClassroom.tsx at it. Set `HALL_MODEL.scale`, `position` and `top` in layout.ts so the hall is about 14.8 m wide with the stage top at STAGE_Y. Adjust WALL_Z, SCREEN, SLIDE and BOARD to the new wall, and the camera presets and zoom limits (MIN_DIST, MAX_DIST, PAN_BOX at the top of BabylonClassroom.tsx). To make a lighter copy of a heavy model: `npx @gltf-transform/cli simplify in.glb out.glb --ratio 0.6 --error 0.0003`.",
  "**New presenter:** use a rigged character with an animation clip as a **GLB** (Mixamo, Ready Player Me or Blender). Babylon.js has no FBX loader, so convert an FBX first; this project's presenter was converted with three.js's FBXLoader and GLTFExporter and then compressed with gltf-transform (`optimize --compress quantize`). Replace `public/models/presenter.glb` and tune `PRESENTER.height`, `x`, `z`, `rotY`.",
  "**Remove the presenter:** delete the `loadPresenter()` call in BabylonClassroom.tsx (and `<Presenter />` in scene/ClassroomCanvas.tsx for the fallback).",
  "**How the presenter works:** skeletal animation (a mesh, bones and a clip), not physics. The clip speed follows the narration and eases to a stop when paused.",
  "**Rendering quality:** Low draws without shadows, glow or anti-aliasing at 30 frames a second; Medium adds shadows, FXAA and glow; High adds 4x anti-aliasing, bigger shadows and fill lights. Weak devices and touch screens are capped automatically, the pixel density steps down if the frame rate drops, and nothing is drawn while the tab is hidden. A still, paused classroom redraws only about 8 times a second.")

H2("4.3 Small screens")
P("One codebase adapts by width, orientation and input. All responsive rules are in `src/styles/90-responsive.css`. Key breakpoints: navigation collapses to a menu below 1140 px; the classroom becomes a portrait player below 620 px and a slim landscape layout on short screens; touch screens get larger controls; reduced-motion stops animation. To hide something on small screens, add a rule such as `@media (max-width: 560px) { .my-banner { display: none; } }`. Test at 320, 375, 414, 768 and 1024 px, portrait and landscape.")
PB()

# ------------------------------ PART 5
H1("Part 5. Help")
H2("5.1 Troubleshooting")
T([["Problem", "What to do"],
   ["Install says Python or Node is missing", "Run the installer again (it installs them). On Windows close the window and run again so the new PATH is seen."],
   ["\"Port 3000 or 8000 is already in use\"", "Close the other copy, or start with --web-port or --api-port."],
   ["Lessons say no AI provider is set up", "Add GEMINI_API_KEY (or another) to .env and restart the API."],
   ["\"No module named ...\" from the backend", "Run backend/.venv/Scripts/python -m pip install -r backend/requirements.txt (Linux or Mac: .venv/bin/python)."],
   ["Classroom is black", "Check WebGL (get.webgl.org), update the graphics driver, try another browser, set 3D quality to Low. The 2D lesson view has the same lesson."],
   ["Changed a NEXT_PUBLIC_ value, nothing happened", "Rebuild the web app; these are fixed at build time."],
   ["Google button missing", "Set both GOOGLE_CLIENT_ID and NEXT_PUBLIC_GOOGLE_CLIENT_ID, add your site as an authorised origin, and rebuild."],
   ["Payments never complete", "Check the Razorpay webhook URL and secret, and that the API is reachable from the internet."],
   ["Emails do not arrive", "Verify the domain with the mail provider (SPF and DKIM) and use a sender on that domain."],
   ["Ads do not show", "Only Free accounts, only after the cookie and ads choice, only with NEXT_PUBLIC_ADSENSE_* set and the site approved."],
   ["Browser stuck on \"Checking your allowance\"", "The rating box (\"How was your last lesson?\") must be answered before the next session; it appears once per unrated session."]], [58, 112])

H2("5.2 Command cheat sheet")
CODE("""python scripts/install.py              install everything          (npm run setup)
python scripts/install.py --doctor     check this computer         (npm run doctor)
python scripts/run.py --open           run API + web, open browser (npm run app)
python scripts/run.py --prod           production build and run    (npm run app:prod)
python scripts/run.py --check          start, run staging checks, stop
npm run typecheck | lint | test        front-end checks
cd backend && python -m unittest discover -s tests -t .      backend tests
python scripts/make_docs.py            rebuild this guide
python scripts/make_mail_images.py     redraw the email header pictures
docker compose up --build -d           production-like stack""")

H2("5.3 Honest limits")
B("AI lessons, marks and podcasts can be wrong. They are study help, not an official record.",
  "The presenter is one looping motion, not lip-sync. The 3D hall is about 6 MB and heavy on old phones; lower the quality or use the lesson view.",
  "Podcast voices and read-aloud lessons use the browser's own voices, so quality varies by device.",
  "Each learner's data is separated by account id in one shared database, not in a separate database per learner.",
  "Limits are counted per server process; with several workers each enforces its own.",
  "Production needs your keys and a staging run. The legal pages are defaults: have them reviewed. Google AdSense needs site approval.")
P("More detail lives in SETUP.md (keys), DEPLOY.md (hosting), BACKEND.md (API and tests), FEATURES.md and .env.example (every setting).")


# ------------------------------------------------------------------------------------------------------------------ render
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def runs(text: str):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            yield part[2:-2], "b"
        elif part.startswith("`"):
            yield part[1:-1], "c"
        else:
            yield part, ""


def build_pdf(path: Path):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Preformatted, Spacer, Table, TableStyle

    ink, acc, soft, line = colors.HexColor("#1b1f2a"), colors.HexColor("#2f5bff"), colors.HexColor("#eef1f8"), colors.HexColor("#cfd5e3")
    ss = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.6, textColor=ink, spaceAfter=5)
    st = {
        "h1": ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=acc, spaceBefore=4, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=16, spaceBefore=10, spaceAfter=4),
        "h3": ParagraphStyle("h3", parent=base, fontName="Helvetica-Bold", fontSize=10, spaceBefore=5, spaceAfter=2),
        "bul": ParagraphStyle("bul", parent=base, leftIndent=13, bulletIndent=3, spaceAfter=2.6),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.5, leading=11, spaceAfter=0),
        "cellb": ParagraphStyle("cellb", parent=base, fontSize=8.5, leading=11, spaceAfter=0, fontName="Helvetica-Bold"),
        "code": ParagraphStyle("code", parent=base, fontName="Courier", fontSize=7.9, leading=10.2, backColor=soft, borderPadding=(5, 6, 5, 6), spaceBefore=3, spaceAfter=9, leftIndent=4, rightIndent=4),
        "note": ParagraphStyle("note", parent=base, backColor=colors.HexColor("#fff6e0"), borderPadding=(6, 7, 6, 7), spaceBefore=4, spaceAfter=9, leftIndent=4, rightIndent=4),
    }

    def mk(t: str) -> str:
        out = []
        for txt, kind in runs(t):
            e = escape(txt)
            out.append(f"<b>{e}</b>" if kind == "b" else f'<font face="Courier" size="8.4">{e}</font>' if kind == "c" else e)
        return "".join(out)

    S = []
    for it in C:
        k = it[0]
        if k == "title":
            S += [Spacer(1, 50), Paragraph(it[1], ParagraphStyle("t", parent=st["h1"], fontSize=36, leading=42)),
                  Paragraph(it[2], ParagraphStyle("s", parent=base, fontSize=12.5, leading=18, spaceAfter=16))]
        elif k in ("h1", "h2", "h3"):
            S.append(Paragraph(mk(it[1]), st[k]))
        elif k == "p":
            S.append(Paragraph(mk(it[1]), base))
        elif k == "note":
            S.append(Paragraph(mk(it[1]), st["note"]))
        elif k == "b":
            S += [Paragraph(mk(x), st["bul"], bulletText="-") for x in it[1]] + [Spacer(1, 4)]
        elif k == "code":
            S.append(Preformatted(it[1], st["code"]))
        elif k == "pb":
            S.append(PageBreak())
        elif k == "t":
            rows, widths = it[1], it[2]
            data = [[Paragraph(mk(str(c)), st["cellb"] if r == 0 else st["cell"]) for c in row] for r, row in enumerate(rows)]
            t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), soft),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            S += [t, Spacer(1, 8)]

    def footer(cv, doc):
        cv.saveState()
        cv.setFont("Helvetica", 8)
        cv.setFillColor(colors.HexColor("#7b8296"))
        cv.drawString(20 * mm, 11 * mm, "WorkLab XC - Complete guide")
        cv.drawRightString(190 * mm, 11 * mm, f"Page {doc.page}")
        cv.restoreState()

    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=20 * mm, title="WorkLab XC - Complete guide", author="WorkLab XC")
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="f")], onPage=footer)])
    doc.build(S)


def build_docx(path: Path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(20)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 11.5)):
        d.styles[name].font.name = "Calibri"
        d.styles[name].font.size = Pt(size)
        d.styles[name].font.color.rgb = RGBColor(0x2F, 0x5B, 0xFF) if name == "Heading 1" else RGBColor(0x1B, 0x1F, 0x2A)

    def shade(cell, hexcolor):
        tc = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcolor)
        tc.append(sh)

    def add_runs(par, text, size=None):
        for txt, kind in runs(text):
            r = par.add_run(txt)
            if kind == "b":
                r.bold = True
            if kind == "c":
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            if size and kind != "c":
                r.font.size = Pt(size)

    for it in C:
        k = it[0]
        if k == "title":
            p = d.add_paragraph()
            r = p.add_run(it[1])
            r.bold = True
            r.font.size = Pt(34)
            r.font.color.rgb = RGBColor(0x2F, 0x5B, 0xFF)
            d.add_paragraph(it[2]).runs[0].font.size = Pt(13)
        elif k == "h1":
            d.add_heading(it[1], level=1)
        elif k == "h2":
            d.add_heading(it[1], level=2)
        elif k == "h3":
            d.add_heading(it[1], level=3)
        elif k == "p":
            add_runs(d.add_paragraph(), it[1])
        elif k == "note":
            p = d.add_paragraph()
            add_runs(p, "Note: " + it[1])
            for r in p.runs:
                r.italic = True
            pPr = p._p.get_or_add_pPr()
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:fill"), "FFF6E0")
            pPr.append(sh)
        elif k == "b":
            for x in it[1]:
                add_runs(d.add_paragraph(style="List Bullet"), x)
        elif k == "code":
            for ln in it[1].split("\n"):
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln or " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                pPr = p._p.get_or_add_pPr()
                sh = OxmlElement("w:shd")
                sh.set(qn("w:val"), "clear")
                sh.set(qn("w:fill"), "EEF1F8")
                pPr.append(sh)
            d.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == "pb":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "t":
            rows, widths = it[1], it[2]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, c in enumerate(row):
                    cell = t.cell(ri, ci)
                    cell.width = Mm(widths[ci])
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], str(c), 9.5)
                    if ri == 0:
                        shade(cell, "EEF1F8")
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            d.add_paragraph()
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_pdf(OUT / "WorkLab-XC-Guide.pdf")
    build_docx(OUT / "WorkLab-XC-Guide.docx")
    print("wrote", OUT / "WorkLab-XC-Guide.pdf", "and", OUT / "WorkLab-XC-Guide.docx")
